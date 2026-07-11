# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import sys, os

sys.path.insert(0, r'C:\Users\mcj\.codex\plugins\cache\openai-primary-runtime\documents\26.630.12135\skills\documents\scripts')
# table_geometry not used

OUTPUT = r'C:\Users\mcj\Desktop\item1\resume.docx'
doc = Document()

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

def srf(run, cn='微软雅黑', en='Calibri', size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = en
    elem = run._element
    rPr = elem.find(qn('w:rPr'))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        elem.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}></w:rFonts>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    if color:
        run.font.color.rgb = color

def sps(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line

def add_centered(doc, text, size=11, bold=False, color=None, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sps(p, before=before, after=after, line=1.15)
    run = p.add_run(text)
    srf(run, size=size, bold=bold, color=color)
    return p

def add_section_header(doc, text):
    p = doc.add_paragraph()
    sps(p, before=16, after=4, line=1.15)
    run = p.add_run(text)
    srf(run, size=14, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    pPr = p._element.find(qn('w:pPr'))
    if pPr is None:
        pPr = parse_xml(f'<w:pPr {nsdecls("w")}></w:pPr>')
        p._element.insert(0, pPr)
    bdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        '  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="1F3A5F"/>'
        '</w:pBdr>'
    )
    pPr.append(bdr)
    return p

def add_body(doc, text, size=10.5, bold=False, before=0, after=2):
    p = doc.add_paragraph()
    sps(p, before=before, after=after, line=1.25)
    run = p.add_run(text)
    srf(run, size=size, bold=bold)
    return p

def add_left_right(doc, left_text, right_text, size=10.5, left_bold=False, right_bold=False, before=0):
    p = doc.add_paragraph()
    sps(p, before=before, after=1, line=1.25)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
    run = p.add_run(left_text)
    srf(run, size=size, bold=left_bold)
    p.add_run('\t')
    run = p.add_run(right_text)
    srf(run, size=size, bold=right_bold, color=RGBColor(0x55, 0x55, 0x55))
    return p

def add_award_item(doc, name, level, desc):
    p = doc.add_paragraph()
    sps(p, before=0, after=0, line=1.25)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.25)
    run = p.add_run('\u25a0 ' + name + '   ')
    srf(run, size=10.5, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    run = p.add_run('[' + level + ']')
    srf(run, size=10, bold=True, color=RGBColor(0xC0, 0x39, 0x2B))
    if desc:
        add_body(doc, desc, size=9.5, before=0, after=2)

def add_bullet(doc, text, size=9.5, indent=0.4):
    p = doc.add_paragraph()
    sps(p, before=0, after=0, line=1.20)
    pf = p.paragraph_format
    pf.left_indent = Inches(indent)
    run = p.add_run('\u2022 ' + text)
    srf(run, size=size)
    return p

# ═══ CONTENT ═══

add_centered(doc, '\u7b80 \u5386', size=20, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F), before=0, after=2)

p = doc.add_paragraph()
sps(p, before=2, after=6, line=1.0)
pPr = p._element.find(qn('w:pPr'))
if pPr is None:
    pPr = parse_xml(f'<w:pPr {nsdecls("w")}></w:pPr>')
    p._element.insert(0, pPr)
bdr = parse_xml(
    f'<w:pBdr {nsdecls("w")}>'
    '  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="1F3A5F"/>'
    '</w:pBdr>'
)
pPr.append(bdr)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sps(p, before=2, after=8, line=1.25)
info_items = ['\u5f20 \u4e09', '\u6e56\u5357\u8b66\u5bdf\u5b66\u9662', '\u6570\u636e\u79d1\u5b66\u4e0e\u5927\u6570\u636e\u6280\u672f', '\u672c\u79d1', '157XXXX1234', 'zhangsan@qq.com']
for i, item in enumerate(info_items):
    run = p.add_run(item)
    srf(run, size=10)
    if i < len(info_items) - 1:
        run = p.add_run('  |  ')
        srf(run, size=9, color=RGBColor(0x99, 0x99, 0x99))

# ═══ Education ═══
add_section_header(doc, '\u6559\u80b2\u80cc\u666f')
add_left_right(doc, '\u6e56\u5357\u8b66\u5bdf\u5b66\u9662', '2023.09 \u2013 2027.06', size=10.5, left_bold=True)
add_body(doc, '\u6570\u636e\u79d1\u5b66\u4e0e\u5927\u6570\u636e\u6280\u672f \u4e13\u4e1a \u00b7 \u672c\u79d1', size=10.5, before=0, after=0)
add_body(doc, '\u4e3b\u4fee\u8bfe\u7a0b\uff1a\u6570\u636e\u7ed3\u6784\u4e0e\u7b97\u6cd5\u3001\u6570\u636e\u5e93\u539f\u7406\u3001Python\u6570\u636e\u5206\u6790\u3001\u673a\u5668\u5b66\u4e60\u3001\u6570\u636e\u53ef\u89c6\u5316\u3001\u5927\u6570\u636e\u6280\u672f\u57fa\u7840', size=9.5, before=1, after=0)

# ═══ Awards ═══
add_section_header(doc, '\u7ade\u8d5b\u83b7\u5956')

awards = [
    ('\u84dd\u6865\u676f\u5168\u56fd\u8f6f\u4ef6\u548c\u4fe1\u606f\u6280\u672f\u4e13\u4e1a\u4eba\u624d\u5927\u8d5b', '\u7701\u7ea7\u4e8c\u7b49\u5956', 'C/C++ \u7ec4\u522b\uff0c\u7efc\u5408\u8003\u5bdf\u7b97\u6cd5\u8bbe\u8ba1\u4e0e\u7f16\u7a0b\u80fd\u529b\uff0c\u6db5\u76d6\u6570\u636e\u7ed3\u6784\u3001\u52a8\u6001\u89c4\u5212\u3001\u56fe\u8bba\u7b49\u6838\u5fc3\u7b97\u6cd5\u6a21\u5757'),
    ('\u4f20\u667a\u676f\u5168\u56fd\u5927\u5b66\u751fIT\u6280\u80fd\u5927\u8d5b', '\u5168\u56fd\u4e09\u7b49\u5956', '\u7efc\u5408\u8fd0\u7528\u8ba1\u7b97\u673a\u57fa\u7840\u77e5\u8bc6\u4e0e\u5de5\u7a0b\u5b9e\u8df5\u80fd\u529b\u5b8c\u6210\u8d5b\u9898\uff0c\u4f53\u73b0\u5e94\u7528\u5f00\u53d1\u4e0e\u95ee\u9898\u89e3\u51b3\u80fd\u529b'),
    ('\u7b2c\u5341\u4e94\u5c4aAPMCM\u4e9a\u592a\u5730\u533a\u5927\u5b66\u751f\u6570\u5b66\u5efa\u6a21\u7ade\u8d5b', '\u4e09\u7b49\u5956', '\u9488\u5bf9\u5b9e\u9645\u95ee\u9898\u5efa\u7acb\u6570\u5b66\u6a21\u578b\uff0c\u4f7f\u7528 Python \u5b8c\u6210\u6570\u636e\u6e05\u6d17\u3001\u7279\u5f81\u5de5\u7a0b\u4e0e\u6a21\u578b\u6c42\u89e3\uff0c\u64b0\u5199\u5168\u82f1\u6587\u7ade\u8d5b\u8bba\u6587'),
]
for name, level, desc in awards:
    add_award_item(doc, name, level, desc)

# ═══ Skills ═══
add_section_header(doc, '\u4e13\u4e1a\u6280\u80fd')

skills = [
    '\u7f16\u7a0b\u8bed\u8a00\uff1a\u719f\u6089 Python\uff0c\u638c\u63e1 C/C++\u3001SQL\uff0c\u5177\u5907\u624e\u5b9e\u7684\u7b97\u6cd5\u57fa\u7840\u4e0e\u7f16\u7a0b\u89c4\u8303\u610f\u8bc6',
    '\u6570\u636e\u5206\u6790\uff1a\u719f\u7ec3\u4f7f\u7528 Pandas\u3001NumPy \u8fdb\u884c\u6570\u636e\u6e05\u6d17\u4e0e\u5904\u7406\uff0c\u8fd0\u7528 Matplotlib\u3001Seaborn \u5b8c\u6210\u53ef\u89c6\u5316',
    '\u673a\u5668\u5b66\u4e60\uff1a\u4e86\u89e3 scikit-learn \u6846\u67b6\uff0c\u638c\u63e1\u56de\u5f52\u3001\u5206\u7c7b\u3001\u805a\u7c7b\u7b49\u5e38\u7528\u6a21\u578b\u7684\u539f\u7406\u4e0e\u8c03\u7528\u65b9\u6cd5',
    '\u5f00\u53d1\u5de5\u5177\uff1a\u719f\u7ec3\u4f7f\u7528 Git \u7248\u672c\u7ba1\u7406\u3001Jupyter Notebook \u4ea4\u4e92\u5f00\u53d1\uff0c\u638c\u63e1 Linux \u57fa\u672c\u64cd\u4f5c\u547d\u4ee4',
    '\u7efc\u5408\u7d20\u517b\uff1a\u5177\u5907\u6570\u5b66\u5efa\u6a21\u4e0e\u82f1\u6587\u79d1\u6280\u6587\u732e\u9605\u8bfb\u80fd\u529b\uff0c\u6709\u4e0d\u9519\u7684\u6587\u6863\u5199\u4f5c\u548c\u56e2\u961f\u534f\u4f5c\u80fd\u529b',
]
for s in skills:
    p = doc.add_paragraph()
    sps(p, before=0, after=1, line=1.25)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.25)
    run = p.add_run('\u2022 ' + s)
    srf(run, size=10)

# ═══ Projects ═══
add_section_header(doc, '\u9879\u76ee\u7ecf\u5386')

add_left_right(doc, '\u516c\u5171\u5b89\u5168\u6570\u636e\u5206\u6790\u4e0e\u53ef\u89c6\u5316', '2025.03 \u2013 2025.06', size=10.5, left_bold=True)
add_body(doc, '\u6280\u672f\u6808\uff1aPython / Pandas / Matplotlib / GeoPandas', size=9.5, before=1, after=1)
for b in ['\u57fa\u4e8e\u67d0\u57ce\u5e02\u516c\u5f00\u72af\u7f6a\u6570\u636e\u96c6\uff0c\u5bf9 5 \u4e07\u4f59\u6761\u8bb0\u5f55\u8fdb\u884c\u6570\u636e\u6e05\u6d17\u3001\u7279\u5f81\u63d0\u53d6\u4e0e\u65f6\u95f4\u5e8f\u5217\u5206\u6790',
          '\u4f7f\u7528 Matplotlib \u4e0e GeoPandas \u7ed8\u5236\u6848\u53d1\u70ed\u529b\u56fe\u4e0e\u8d8b\u52bf\u6298\u7ebf\u56fe\uff0c\u6709\u6548\u6316\u6398\u72af\u7f6a\u9ad8\u53d1\u65f6\u6bb5\u4e0e\u533a\u57df\u5206\u5e03\u89c4\u5f8b',
          '\u64b0\u5199\u6570\u636e\u5206\u6790\u62a5\u544a\u5e76\u5236\u4f5c\u53ef\u89c6\u5316\u5c55\u793a\uff0c\u83b7\u9662\u7ea7\u6570\u636e\u6316\u6398\u7ade\u8d5b\u4f18\u79c0\u4f5c\u54c1\u5956']:
    add_bullet(doc, b)

add_left_right(doc, '\u7535\u5546\u7528\u6237\u8d2d\u4e70\u884c\u4e3a\u9884\u6d4b', '2024.10 \u2013 2024.12', size=10.5, left_bold=True, before=6)
add_body(doc, '\u6280\u672f\u6808\uff1aPython / scikit-learn / Pandas / Jupyter', size=9.5, before=1, after=1)
for b in ['\u4f7f\u7528 Kaggle Online Retail \u6570\u636e\u96c6\uff0c\u5b8c\u6210\u7f3a\u5931\u503c\u5904\u7406\u3001\u7279\u5f81\u5de5\u7a0b\u4e0e\u7528\u6237\u753b\u50cf\u6784\u5efa',
          '\u57fa\u4e8e\u903b\u8f91\u56de\u5f52\u4e0e\u968f\u673a\u68ee\u6797\u6784\u5efa\u8d2d\u4e70\u884c\u4e3a\u9884\u6d4b\u6a21\u578b\uff0c\u51c6\u786e\u7387\u8fbe 87%\uff0cAUC \u4e3a 0.92',
          '\u901a\u8fc7\u7279\u5f81\u91cd\u8981\u6027\u5206\u6790\u8bc6\u522b\u5f71\u54cd\u8d2d\u4e70\u51b3\u7b56\u7684\u5173\u952e\u56e0\u5b50\uff0c\u5e76\u7ed9\u51fa\u8425\u9500\u7b56\u7565\u4f18\u5316\u5efa\u8bae']:
    add_bullet(doc, b)

add_left_right(doc, '\u84dd\u6865\u676f\u7b97\u6cd5\u7ade\u8d5b\u8bad\u7ec3\u9898\u5e93', '2024.03 \u2013 2024.06', size=10.5, left_bold=True, before=6)
add_body(doc, '\u6280\u672f\u6808\uff1aC/C++ / \u6570\u636e\u7ed3\u6784 / \u7b97\u6cd5\u8bbe\u8ba1', size=9.5, before=1, after=1)
for b in ['\u7cfb\u7edf\u6574\u7406 60+ \u9053\u7ade\u8d5b\u771f\u9898\uff0c\u8986\u76d6\u52a8\u6001\u89c4\u5212\u3001\u56fe\u8bba\u3001\u5b57\u7b26\u4e32\u5904\u7406\u3001\u8d2a\u5fc3\u7b97\u6cd5\u7b49\u9ad8\u9891\u8003\u70b9',
          '\u4e3a\u6bcf\u9053\u9898\u7f16\u5199\u591a\u89e3\u6cd5\u5206\u6790\u4e0e\u590d\u6742\u5ea6\u5bf9\u6bd4\u6587\u6863\uff0c\u7528\u4e8e\u5907\u8d5b\u5c0f\u7ec4\u4ea4\u6d41\u4e0e\u81ea\u4e3b\u5b66\u4e60',
          '\u5728\u84dd\u6865\u676f\u7701\u8d5b\u4e2d\u83b7\u4e8c\u7b49\u5956\uff0c\u7b97\u6cd5\u7efc\u5408\u6392\u540d\u4f4d\u5217\u8d5b\u70b9\u524d 15%']:
    add_bullet(doc, b)

doc.save(OUTPUT)
print('Done: ' + OUTPUT)

